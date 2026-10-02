# -*- coding: utf-8 -*-
"""学生成绩单导出探针（2026-10-02，W1.8）。

判据（全部带负向对照）：
  1. 导出本人成绩单 → HTTP 200 + xlsx 内容类型 + 响应体真是 xlsx（zip 头 / openpyxl 能打开）
  2. 工作簿行数 = 同一学生的 `/students/scores/{uid}` 条数（**同一份数据，两个端点对得上**）
  3. 学期列是可读区间（"2023-2024-1"），不是裸的全局序号（"1"）
  4. 不带身份头 / 只带 X-User-Id → 403（GatewayAuthFilter）
  5. 拿别人的 id 导出 → 403 且提示"本人"
  6. 不存在的用户 → 404

⚠️ 只读，不写任何表。
⚠️ 中文一律经 API/工作簿回读判定，不经 CLI stdout（CP936 乱码）。
"""
import io
import json
import sys
import urllib.request
import urllib.error

import openpyxl

BASE = "http://localhost:%s" % (sys.argv[1] if len(sys.argv) > 1 else "8179")
ME = 27            # 司明（users.id）
OTHER = 30
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

failures = []


def check(name, ok, detail=""):
    print(("[PASS] " if ok else "[FAIL] ") + name + (("  " + detail) if detail else ""))
    if not ok:
        failures.append(name)


def get(path, headers=None):
    req = urllib.request.Request(BASE + path, headers=headers or {})
    try:
        with urllib.request.urlopen(req, timeout=20) as resp:
            return resp.status, resp.headers.get("Content-Type", ""), resp.read()
    except urllib.error.HTTPError as e:
        return e.code, e.headers.get("Content-Type", ""), e.read()


def as_json(raw):
    try:
        return json.loads(raw.decode("utf-8"))
    except Exception:
        return {}


def auth(uid, role="1"):
    return {"X-User-Id": str(uid), "X-Role": role}


# --- 基线：接口侧的成绩条数 ---
status, _, raw = get("/students/scores/%d" % ME, auth(ME))
body = as_json(raw)
scores = body.get("data") if isinstance(body.get("data"), list) else []
print("baseline: /students/scores/%d -> http=%s, 条数=%d" % (ME, status, len(scores)))
check("基线接口取到成绩", status == 200 and len(scores) > 0, "http=%s n=%d" % (status, len(scores)))

# --- 1. 导出本人成绩单 ---
status, ctype, raw = get("/students/export/scores/%d/excel" % ME, auth(ME))
check("导出本人成绩单 HTTP 200", status == 200, "http=%s" % status)
check("内容类型是 xlsx", ctype.startswith(XLSX_TYPE), ctype)
check("响应体是 zip 容器（xlsx）而非 JSON",
      len(raw) > 2 and raw[0] == 0x50 and raw[1] == 0x4B,
      "前两字节=%s 长度=%d" % (list(raw[:2]), len(raw)))

wb = None
try:
    wb = openpyxl.load_workbook(io.BytesIO(raw))
except Exception as e:
    check("openpyxl 能打开工作簿", False, str(e))

if wb is not None:
    check("openpyxl 能打开工作簿", True)
    ws = wb[wb.sheetnames[0]]
    header = [c.value for c in ws[2]]
    check("表头正确",
          header[:3] == ["学期", "课程名称", "学分"],
          json.dumps(header, ensure_ascii=False))

    data_rows = list(ws.iter_rows(min_row=3, values_only=True))
    data_rows = [r for r in data_rows if any(v not in (None, "") for v in r)]
    check("数据行数 = 接口侧条数", len(data_rows) == len(scores),
          "工作簿=%d 接口=%d" % (len(data_rows), len(scores)))

    semesters = [r[0] for r in data_rows]
    check("学期列是区间字符串（不是裸序号）",
          all(isinstance(s, str) and s.count("-") == 2 and s[:4].isdigit() for s in semesters),
          json.dumps(semesters, ensure_ascii=False))

    names = {r[1] for r in data_rows}
    api_courses = {s.get("courseName") for s in scores if s.get("courseName")}
    check("课程名与接口一致", names == api_courses,
          "差集=%s" % json.dumps(sorted(names ^ api_courses), ensure_ascii=False))

# --- 负向对照 ---
status, _, _ = get("/students/export/scores/%d/excel" % ME)
check("不带身份头 → 403", status == 403, "http=%s" % status)

status, _, _ = get("/students/export/scores/%d/excel" % ME, {"X-User-Id": str(ME)})
check("只带 X-User-Id、缺 X-Role → 403", status == 403, "http=%s" % status)

status, ctype, raw = get("/students/export/scores/%d/excel" % ME, auth(OTHER))
check("导出别人的成绩单 → 403", status == 403, "http=%s" % status)
check("403 的响应体是 JSON 错误（前端才分得清）",
      "json" in ctype.lower() and "本人" in raw.decode("utf-8", "replace"),
      "ctype=%s body=%s" % (ctype, raw.decode("utf-8", "replace")[:80]))

status, _, raw = get("/students/export/scores/99999/excel", auth(99999))
check("不存在的用户 → 404", status == 404, "http=%s" % status)

print()
print("FAILURES: %d" % len(failures))
sys.exit(1 if failures else 0)
